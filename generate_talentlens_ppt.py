from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'TalentLens_Project_Presentation.pptx'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

DARK = RGBColor(9, 30, 52)
NAVY = RGBColor(17, 56, 87)
BLUE = RGBColor(38, 139, 230)
LIGHT = RGBColor(225, 236, 247)
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(98, 196, 255)
GREEN = RGBColor(55, 180, 120)
TEXT_DARK = RGBColor(23, 37, 84)
GRAY = RGBColor(99, 115, 129)


def add_background(slide, color=DARK):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_block(slide, title, subtitle=None, accent=True):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.7), Inches(8.5), Inches(1.0))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE if accent else TEXT_DARK
    p.alignment = PP_ALIGN.LEFT

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.45), Inches(9.2), Inches(0.6))
        p2 = sub_box.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(12)
        p2.font.color.rgb = LIGHT if accent else GRAY


def add_section_header(slide, text, y=0.45):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(y), Inches(12.5), Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.65), Inches(y - 0.08), Inches(8.0), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE


def add_bullet_box(slide, left, top, width, height, bullets, title=None, color=WHITE, text_color=TEXT_DARK, title_size=16, body_size=12):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = WHITE
    shape.line.width = Pt(1)

    if title:
        tbox = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.12), width - Inches(0.28), Inches(0.38))
        tp = tbox.text_frame.paragraphs[0]
        tp.text = title
        tp.font.size = Pt(title_size)
        tp.font.bold = True
        tp.font.color.rgb = text_color

    box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.52), width - Inches(0.35), height - Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.font.size = Pt(body_size)
        p.font.color.rgb = text_color
        p.space_after = Pt(4)


def add_feature_tile(slide, left, top, width, height, title, text, icon='•'):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(18, 52, 79)
    shape.line.color.rgb = BLUE
    shape.line.width = Pt(1.2)
    icon_box = slide.shapes.add_textbox(left + Inches(0.16), top + Inches(0.14), Inches(0.28), Inches(0.3))
    icon_p = icon_box.text_frame.paragraphs[0]
    icon_p.text = icon
    icon_p.font.size = Pt(18)
    icon_p.font.bold = True
    icon_p.font.color.rgb = ACCENT

    tb = slide.shapes.add_textbox(left + Inches(0.52), top + Inches(0.1), width - Inches(0.62), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE

    tb2 = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.58), width - Inches(0.28), height - Inches(0.68))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = text
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = LIGHT
    p2.word_wrap = True


# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.18))
shape.fill.solid(); shape.fill.fore_color.rgb = BLUE; shape.line.fill.background()
logo_path = ROOT / 'frontend' / 'public' / 'icon.png'
if logo_path.exists():
    slide.shapes.add_picture(str(logo_path), Inches(0.8), Inches(0.55), Inches(0.72), Inches(0.72))

add_title_block(slide, 'TalentLens', 'AI-powered resume analytics and job recommendation platform', accent=True)

# Main highlight cards
card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.2), Inches(3.4), Inches(1.8))
card1.fill.solid(); card1.fill.fore_color.rgb = NAVY; card1.line.color.rgb = BLUE
card1.line.width = Pt(1.5)
text1 = card1.text_frame
text1.word_wrap = True
p1 = text1.paragraphs[0]
p1.text = 'Resume-to-Job Fit'
p1.font.size = Pt(20); p1.font.bold = True; p1.font.color.rgb = WHITE
p2 = text1.add_paragraph(); p2.text = 'Semantic analysis beyond keyword matching'; p2.font.size = Pt(12); p2.font.color.rgb = LIGHT

card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.7), Inches(2.2), Inches(3.4), Inches(1.8))
card2.fill.solid(); card2.fill.fore_color.rgb = NAVY; card2.line.color.rgb = BLUE
card2.line.width = Pt(1.5)
text2 = card2.text_frame
text2.word_wrap = True
p1 = text2.paragraphs[0]
p1.text = 'Career Role Match'
p1.font.size = Pt(20); p1.font.bold = True; p1.font.color.rgb = WHITE
p2 = text2.add_paragraph(); p2.text = 'Strong, alternative, and missing-skill recommendations'; p2.font.size = Pt(12); p2.font.color.rgb = LIGHT

card3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(2.2), Inches(3.8), Inches(1.8))
card3.fill.solid(); card3.fill.fore_color.rgb = NAVY; card3.line.color.rgb = BLUE
card3.line.width = Pt(1.5)
text3 = card3.text_frame
text3.word_wrap = True
p1 = text3.paragraphs[0]
p1.text = 'Portfolio + Job Discovery'
p1.font.size = Pt(20); p1.font.bold = True; p1.font.color.rgb = WHITE
p2 = text3.add_paragraph(); p2.text = 'Editable portfolio preview and relevant India-focused opportunities'; p2.font.size = Pt(12); p2.font.color.rgb = LIGHT

# Footer note
note = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5))
pt = note.text_frame.paragraphs[0]
pt.text = 'Final-year major project | Smart career guidance for job seekers'
pt.font.size = Pt(14)
pt.font.color.rgb = LIGHT
pt.alignment = PP_ALIGN.CENTER

# Slide 2: Problem + Motivation
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Problem Statement', y=0.45)

left = Inches(0.7)
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.0), Inches(5.5), Inches(4.9))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(17, 52, 76); box.line.color.rgb = BLUE

text = box.text_frame
text.word_wrap = True
for text_line, size in [
    ('Traditional resume screening often relies on keyword matching.', 18),
    ('This fails to capture:', 16),
    ('• Skill context and usage', 14),
    ('• Career-role suitability', 14),
    ('• Missing skills and gaps', 14),
    ('• Project depth and achievements', 14),
    ('• Semantic alignment between resume and job requirement', 14),
]:
    p = text.paragraphs[0] if not text.paragraphs[0].text else text.add_paragraph()
    p.text = text_line
    p.font.size = Pt(size)
    p.font.color.rgb = WHITE
    p.bullet = text_line.startswith('•')
    p.level = 0

box2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(1.0), Inches(5.8), Inches(4.9))
box2.fill.solid(); box2.fill.fore_color.rgb = RGBColor(28, 72, 93); box2.line.color.rgb = BLUE
text2 = box2.text_frame
for idx, line in enumerate([
    'Why TalentLens?',
    'To help job seekers understand their profile more realistically,',
    'identify useful career directions,',
    'discover suitable opportunities, and improve employability with explainable insights.',
]):
    p = text2.paragraphs[0] if idx == 0 else text2.add_paragraph()
    p.text = line
    p.font.size = Pt(22 if idx == 0 else 15)
    p.font.bold = (idx == 0)
    p.font.color.rgb = WHITE

# Slide 3: Objectives
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Project Objectives', y=0.45)
add_feature_tile(slide, Inches(0.7), Inches(1.1), Inches(3.7), Inches(2.1), '1. Resume Analytics', 'Extract structured resume information and understand skills, experience, education, projects, and strengths.')
add_feature_tile(slide, Inches(4.7), Inches(1.1), Inches(3.7), Inches(2.1), '2. Career Matching', 'Recommend the best-fit and alternative job roles using contextual reasoning rather than simple keyword overlap.')
add_feature_tile(slide, Inches(8.7), Inches(1.1), Inches(3.9), Inches(2.1), '3. Skill Gap Detection', 'Highlight important missing skills and suggest useful resume improvements for better employability.')
add_feature_tile(slide, Inches(0.7), Inches(3.7), Inches(3.7), Inches(2.1), '4. Portfolio Transformation', 'Turn extracted profile data into an editable, professional portfolio preview that can be printed or saved as PDF.')
add_feature_tile(slide, Inches(4.7), Inches(3.7), Inches(3.7), Inches(2.1), '5. Domain-Aware Scoring', 'Apply domain-based skill weighting with bounded, explainable score components for transparent job matching.')
add_feature_tile(slide, Inches(8.7), Inches(3.7), Inches(3.9), Inches(2.1), '6. Evolving Skill Catalogue', 'Allow users to review and accept new skills, which are normalized and saved for future recommendation accuracy.')

# Slide 4: Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Key Features', y=0.45)
features = [
    ('Resume upload', 'Supports PDF and DOCX resume processing.'),
    ('AI-powered analysis', 'Uses Google Gemini for contextual resume reasoning.'),
    ('Career role suggestions', 'Best match and alternative roles are identified dynamically.'),
    ('Skill gap analysis', 'Detects missing skills and resume improvement opportunities.'),
    ('Job discovery', 'Searches and normalizes relevant jobs, primarily India-oriented.'),
    ('Resume-to-job matching', 'Combines baseline scoring with semantic validation.'),
    ('Portfolio preview', 'Editable professional portfolio with browser export options.'),
    ('Secure authentication', 'JWT-based access and MongoDB-backed user data.'),
]
for idx, (title, text) in enumerate(features):
    col = idx % 4
    row = idx // 4
    left = Inches(0.6 + col * 3.15)
    top = Inches(1.1 + row * 2.0)
    add_feature_tile(slide, left, top, Inches(2.8), Inches(1.55), title, text, icon='✓')

# Slide 5: Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'System Architecture', y=0.45)
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.0), Inches(12.0), Inches(4.8))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(18, 52, 79); box.line.color.rgb = BLUE
text = box.text_frame
text.word_wrap = True
arch = [
    'Client / Frontend (React + Vite)',
    '↓',
    'Backend API (Express.js + Node.js)',
    '↓',
    'Services: Resume parsing, Gemini analysis, skill matching, job discovery',
    '↓',
    'Database: MongoDB for users, skills, and persisted catalogue data',
    '↓',
    'External integration: Google Gemini API + Jobicy public jobs API'
]
for i, line in enumerate(arch):
    p = text.paragraphs[0] if i == 0 else text.add_paragraph()
    p.text = line
    p.font.size = Pt(18 if i % 2 == 0 else 20)
    p.font.bold = (i % 2 == 1 or (i in [0, 2, 4, 6]))
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

# Slide 6: Tech Stack
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Technology Stack', y=0.45)
add_bullet_box(slide, Inches(0.7), Inches(1.2), Inches(3.5), Inches(3.8), ['React', 'Vite', 'JavaScript', 'Responsive UI', 'Portfolio export'], title='Frontend', color=RGBColor(17, 56, 87), text_color=WHITE)
add_bullet_box(slide, Inches(4.6), Inches(1.2), Inches(3.5), Inches(3.8), ['Node.js', 'Express.js', 'REST APIs', 'JWT auth', 'Multer uploads'], title='Backend', color=RGBColor(17, 56, 87), text_color=WHITE)
add_bullet_box(slide, Inches(8.5), Inches(1.2), Inches(3.9), Inches(3.8), ['MongoDB', 'Google Gemini API', 'Jobicy API', 'Resume parsing (PDF/DOCX)', 'AI embeddings / scoring'], title='Data & Intelligence', color=RGBColor(17, 56, 87), text_color=WHITE)

# Slide 7: Resume Analysis Pipeline
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Resume Analysis Pipeline', y=0.45)
steps = [
    ('1. Upload', 'User uploads PDF or DOCX resume.'),
    ('2. Extract text', 'Backend reads and parses resume content.'),
    ('3. Normalize', 'Resume information is structured into skills, education, experience, projects, and certifications.'),
    ('4. AI reasoning', 'Gemini interprets context, strengths, role fit, and improvement opportunities.'),
    ('5. Validate', 'Results are checked and normalized before the frontend renders them.'),
    ('6. Recommend', 'Best-fit role, skill gaps, and job opportunities are generated.'),
]
for idx, (title, text) in enumerate(steps):
    left = Inches(0.7 + (idx % 3) * 4.1)
    top = Inches(1.1 + (idx // 3) * 2.25)
    add_feature_tile(slide, left, top, Inches(3.6), Inches(1.7), title, text, icon=str(idx+1))

# Slide 8: Job Recommendation & Matching
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Job Recommendation & Matching', y=0.45)
add_bullet_box(slide, Inches(0.8), Inches(1.2), Inches(5.4), Inches(4.5), ['Collects role, location, and resume context', 'Queries external job provider data', 'Filters and ranks India-relevant opportunities', 'Displays job title, company, source, and URL'], title='Discovery', color=RGBColor(17, 56, 87), text_color=WHITE)
add_bullet_box(slide, Inches(6.7), Inches(1.2), Inches(5.8), Inches(4.5), ['Uses baseline keyword match as a foundation', 'Adds semantic evaluation with Gemini embeddings when available', 'Applies deterministic fallback if embeddings are unavailable', 'Explains matched vs missing skills and score components'], title='Matching Logic', color=RGBColor(17, 56, 87), text_color=WHITE)

# Slide 9: Current Status & Roadmap
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Current Status & Future Roadmap', y=0.45)
add_bullet_box(slide, Inches(0.8), Inches(1.1), Inches(5.0), Inches(4.7), ['User login/register', 'Resume parsing and AI analysis', 'Career role recommendation', 'Skill gap and improvement suggestions', 'Portfolio preview and PDF export', 'India-oriented job discovery'], title='Implemented', color=RGBColor(17, 56, 87), text_color=WHITE)
add_bullet_box(slide, Inches(6.2), Inches(1.1), Inches(5.9), Inches(4.7), ['Recruiter-side functionality', 'Broader job provider coverage', 'Candidate ranking and employer dashboards', 'Portfolio publishing and sharing', 'Improved semantic retrieval and ranking', 'Expanded product maturity and deployment'], title='Planned Enhancements', color=RGBColor(23, 82, 104), text_color=WHITE)

# Slide 10: Team / Closing
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide)
add_section_header(slide, 'Project Summary', y=0.45)

summary = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.5), Inches(2.2))
ps = summary.text_frame.paragraphs[0]
ps.text = 'TalentLens is an AI-powered career assistance platform designed to help job seekers analyze their resumes, identify suitable job roles, discover relevant opportunities, and improve their employability with transparent, explainable insights.'
ps.font.size = Pt(22)
ps.font.color.rgb = WHITE
ps.alignment = PP_ALIGN.LEFT

team = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(5.4), Inches(1.4))
pt = team.text_frame.paragraphs[0]
pt.text = 'Team: Final-year major project team\nFocus: Smart, AI-assisted career guidance for students and job seekers'
pt.font.size = Pt(18)
pt.font.color.rgb = LIGHT

thanks = slide.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(11.5), Inches(0.7))
pt2 = thanks.text_frame.paragraphs[0]
pt2.text = 'Thank You'
pt2.font.size = Pt(26)
pt2.font.bold = True
pt2.font.color.rgb = ACCENT
pt2.alignment = PP_ALIGN.CENTER

prs.save(OUTPUT)
print(f'Created presentation: {OUTPUT}')
